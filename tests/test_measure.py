"""Tests for the measurement code in measure/ (M3, M4, M5, M15).

The measurements report counts such as "0 problems". These tests make sure such a count is
earned: each checker finds planted mistakes on the same path it uses for real drafts.
"""
import ast
import copy
import datetime as dt
import io
import json
import pathlib
import re
import shutil
import warnings

import pytest

from measure import datachecks, faults, provenance, publisher
from measure.controls import run_controls
from measure.independent import check_commentary, run_of_changes
from scripts import measure_offline

warnings.filterwarnings("ignore", category=DeprecationWarning)

ROOT = pathlib.Path(__file__).resolve().parent.parent
SNAPSHOT = json.loads((ROOT / "data" / "snapshots" / "dgbf_2026-08.json").read_text(encoding="utf-8"))
needs_node = pytest.mark.skipif(shutil.which("node") is None, reason="Node.js renders the margin notes")


@pytest.fixture(scope="module")
def result():
    return measure_offline.demo_result()


# ---- independence ----------------------------------------------------------------------

@pytest.mark.parametrize("name", ["independent.py", "publisher.py", "controls.py", "datachecks.py"])
def test_independent_checkers_import_nothing_from_the_product(name):
    tree = ast.parse((ROOT / "measure" / name).read_text(encoding="utf-8"))
    imported = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imported |= {a.name.split(".")[0] for a in node.names}
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module.split(".")[0])
    assert "app" not in imported


# ---- M3: independent recomputation -----------------------------------------------------

def test_every_figure_in_the_demo_draft_matches_the_recomputation(result):
    r = check_commentary(result["commentary"], SNAPSHOT)
    assert r["numbers"] > 0 and r["text_with_figures"] > 0
    assert r["problems"] == 0, r["failed_tokens"]
    assert r["uncovered"] == [] and r["digits_outside_tokens"] == 0 and r["unrendered_placeholders"] == 0


def test_planted_mistakes_are_found_and_correct_renderings_accepted(result):
    report = run_controls(result["commentary"], SNAPSHOT)
    wrong = [c for c in report["controls"] if (c["expect"] == "detected") != c["detected"]]
    assert report["all_as_expected"], wrong
    assert sum(c["expect"] == "detected" and c["applicable"] for c in report["controls"]) >= 12
    assert sum(c["expect"] == "accepted" for c in report["controls"]) >= 2


@pytest.mark.parametrize("values,expected", [
    ([1, 2, 3, 4], (3, 1, False)),          # the run reaches the start of the data
    ([5, 1, 2, 3], (2, 1, True)),
    ([1, 2, 1, 0], (2, -1, True)),
    ([1, 1], (0, 0, True)),
])
def test_run_of_changes(values, expected):
    assert run_of_changes(values) == expected


def test_japanese_numeral_scan_ignores_ordinary_words():
    commentary = [{"key": "outlook", "sentences": [{"ref": "outlook.1", "en": [],
                   "ja": [{"type": "text", "text": "一方、動向を十分に注視してまいります。"}]}]}]
    assert check_commentary(commentary, SNAPSHOT)["number_words_outside_tokens"] == []
    commentary[0]["sentences"][0]["ja"][0]["text"] = "約三分の一を占めました。"
    assert check_commentary(commentary, SNAPSHOT)["number_words_outside_tokens"]


# ---- M3 step 0: publisher files ----------------------------------------------------------

def _product_xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Past performance is not a reliable indicator of future performance."])
    ws.append(["As of** ", "Ticker", "Name", "Total Returns as of Date", "Total Returns (Cumulative)", None, None])
    ws.append([None, None, None, None, "1 Month", "QTD", "YTD"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _nav_xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Fund Name:", "Example"])
    ws.append([None])
    ws.append(["Date", "NAV", "Shares Outstanding"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_publisher_parsers_read_the_published_layouts():
    product = publisher.parse_product_data(_product_xlsx([
        ["Sep 22 2026", "SPY®", "Example", "Aug 31 2026", "2.71%", "-", "13.04%"],
        ["Sep 22 2026", "SPAB", "Example", "Aug 31 2026", "0.39%", "-", "-0.23%"],
    ]))
    assert product["SPY"] == {"returns_as_of": "2026-08-31", "tr_1m": 2.71, "tr_ytd": 13.04}
    assert product["SPAB"]["tr_ytd"] == -0.23
    navs = publisher.parse_nav_history(_nav_xlsx([["31-Aug-2026", 25.098868, 1.0],
                                                  [dt.datetime(2026, 7, 31), 25.09106, 1.0]]))
    assert navs == {"2026-08-31": 25.098868, "2026-07-31": 25.09106}
    table = publisher.parse_sdmx_csv("\ufeffDATAFLOW,REF_AREA,TIME_PERIOD,OBS_VALUE\r\nX,G7,2026-08,100.8786\r\nX,G7,2026-07,\r\n")
    assert table == {("G7", "2026-08"): 100.8786}


def test_publisher_parser_refuses_an_unknown_layout():
    with pytest.raises(ValueError):
        publisher.parse_product_data(_nav_xlsx([["31-Aug-2026", 1.0, 1.0]]))


def test_comparison_finds_a_planted_difference():
    product = {"SPY": {"returns_as_of": "2026-08-31", "tr_1m": 2.71, "tr_ytd": 13.04}}
    snap = copy.deepcopy(SNAPSHOT)
    snap["etfs"] = {"SPY": snap["etfs"]["SPY"]}
    navs = {"SPY": {snap["etfs"]["SPY"]["nav_prev_month_end"]["date"]: snap["etfs"]["SPY"]["nav_prev_month_end"]["nav"],
                    snap["etfs"]["SPY"]["nav_month_end"]["date"]: snap["etfs"]["SPY"]["nav_month_end"]["nav"]}}
    cli = {(a, p): v for a, e in snap["macro"]["cli"].items() for p, v in e["series"]}
    rates = {(a, p): v for a, e in snap["macro"]["lt_rate"].items() for p, v in e["series"]}
    rows = publisher.compare(snap, product, navs, cli, rates)
    assert rows and all(r["equal"] for r in rows)
    snap["etfs"]["SPY"]["tr_1m"] = 2.72
    snap["macro"]["cli"]["G7"]["series"][-1][1] += 0.0001
    different = [r["item"] for r in publisher.compare(snap, product, navs, cli, rates) if not r["equal"]]
    assert different == ["SPY 1-month total return, %", "OECD leading indicator G7 2026-08"]


# ---- M4: provenance ----------------------------------------------------------------------

def test_every_figure_traces_to_a_source(result):
    r = provenance.check_provenance(result, ROOT)
    assert r["figure_tokens"] > 0
    assert r["figure_tokens_traced"] == r["figure_tokens"], r["failures"]
    assert "State Street Global Advisors" in r["publishers_reached"] and "OECD" in r["publishers_reached"]


@needs_node
def test_every_figure_has_a_complete_margin_note(result):
    r = provenance.check_provenance(result, ROOT)
    assert r["figure_tokens_note_complete"] == r["figure_tokens"], r["failures"]


def test_tracing_finds_planted_gaps(result):
    sheet = copy.deepcopy(result["fact_sheet"])
    facts, sources = sheet["facts"], {s["key"]: s for s in sheet["sources"]}
    assert provenance.trace("fund.ret_1m", facts, sources)[0] == []
    facts["fund.contrib.SPY"]["derivation"] = None
    assert provenance.trace("fund.ret_1m", facts, sources)[0]
    facts["etf.SPDW.ret_1m"]["source"] = "unknown_vendor"
    assert any("unknown source" in p for p in provenance.trace("fund.contrib.SPDW", facts, sources)[0])
    facts["fund.contrib.SPEM"]["inputs"] = ["etf.SPEM.no_such_fact"]
    assert any("does not exist" in p for p in provenance.trace("fund.contrib.SPEM", facts, sources)[0])


def test_note_check_finds_planted_gaps(result):
    sheet = result["fact_sheet"]
    facts, sources = sheet["facts"], {s["key"]: s for s in sheet["sources"]}
    good = {"value": "3.23%", "body": "Look-through weights ..., as of September 17, 2026. State Street Global Advisors."}
    assert provenance.note_problems(facts["lookthrough.1.weight"], good, "en", sheet, sources) == []
    no_date = {"value": "3.23%", "body": "Look-through weights. State Street Global Advisors."}
    assert provenance.note_problems(facts["lookthrough.1.weight"], no_date, "en", sheet, sources)
    no_publisher = {"value": "3.23%", "body": "Look-through weights, as of September 17, 2026."}
    assert provenance.note_problems(facts["lookthrough.1.weight"], no_publisher, "en", sheet, sources)
    no_calc = {"value": "+1.08%", "body": "寄与度。"}
    assert provenance.note_problems(facts["fund.contrib.SPY"], no_calc, "ja", sheet, sources)


# ---- M15: data checks --------------------------------------------------------------------

def test_data_checks_are_recomputed_and_the_identity_is_not_counted(result):
    r = datachecks.check_datachecks(result, SNAPSHOT)
    assert r["recomputed_agree"] == r["checks"]
    kinds = {row["name"]: row["kind"] for row in r["rows"]}
    assert kinds["Contributions add up to the fund return"] == "identity"
    assert r["verifications"] == r["checks"] - 1
    assert "unclassified" not in r["by_kind"]


def test_data_check_recomputation_finds_a_planted_difference(result):
    snap = copy.deepcopy(SNAPSHOT)
    snap["etfs"]["SPY"]["nav_month_end"]["nav"] += 0.5
    r = datachecks.check_datachecks(result, snap)
    row = next(x for x in r["rows"] if x["name"] == "SPY: published return vs NAV change")
    assert row["agrees"] is False


# ---- M5: fault catalog -------------------------------------------------------------------

def _literals(path: pathlib.Path) -> set[str]:
    tree = ast.parse(path.read_text(encoding="utf-8"))
    return {n.value for n in ast.walk(tree) if isinstance(n, ast.Constant) and isinstance(n.value, str)}


def test_catalog_ids_are_unique_and_complete():
    cat = faults.load_catalog()
    cases = cat["faults"] + cat["legit"]
    ids = [c["id"] for c in cases]
    assert len(ids) == len(set(ids))
    for c in cat["faults"]:
        assert c["family"] in cat["families"] and c["severity"] in ("fatal", "caution")
        assert c["family"] in faults.TARGET_CODES
    assert all(c["severity"] == "accept" for c in cat["legit"])
    assert set(cat["families"]) == {c["family"] for c in cat["faults"]}


def test_catalog_is_not_copied_from_the_unit_tests_or_the_demo_mistakes():
    from app.llm import demo_draft
    known = _literals(ROOT / "tests" / "test_rules.py")
    flawed = demo_draft(set(faults.build_fact_sheet(faults.load_snapshot(ROOT / "data" / "snapshots", "DGBF", "2026-08")).facts), flawed=True)
    known |= {s[lang] for sec in flawed["sections"] for s in sec["sentences"] for lang in ("en", "ja")}
    cat = faults.load_catalog()
    texts = {e[lang] for c in cat["faults"] + cat["legit"] for e in c.get("edits", []) for lang in ("en", "ja") if lang in e}
    copied = {t for t in texts if t in known}
    assert copied == set()


def test_catalog_edits_point_at_existing_sentences():
    from app.llm import demo_draft
    sheet = faults.build_fact_sheet(faults.load_snapshot(ROOT / "data" / "snapshots", "DGBF", "2026-08"))
    base = demo_draft(set(sheet.facts))
    refs = {f"{s['key']}.{i}" for s in base["sections"] for i in range(len(s["sentences"]))}
    cat = faults.load_catalog()
    for c in cat["faults"] + cat["legit"]:
        for e in c.get("edits", []):
            assert e["at"] in refs, c["id"]
        if "drop" in c:
            assert c["drop"] in refs


def test_catalog_file_matches_its_builder(tmp_path):
    source = (ROOT / "measure" / "build_catalog.py").read_text(encoding="utf-8")
    source = source.replace('pathlib.Path(__file__).resolve().parent / "fault_catalog.json"',
                            f'pathlib.Path(r"{tmp_path}") / "fault_catalog.json"')
    exec(compile(source, "build_catalog.py", "exec"), {"__name__": "catalog_builder"})
    assert (tmp_path / "fault_catalog.json").read_bytes() == faults.CATALOG.read_bytes()


def test_fault_run_is_deterministic_and_the_base_draft_passes():
    first = faults.run(ROOT / "data" / "snapshots")
    second = faults.run(ROOT / "data" / "snapshots")
    assert first == second
    assert first["base_draft_issues"] == []
    assert re.fullmatch(r"[0-9a-f]{32}", first["catalog_md5"])
    total = first["faults"]["all"]
    assert total["cases"] == total["blocked_target"] + total["blocked_other"] + total["flagged"] + total["passed"]


def test_fault_injection_sees_a_blocked_and_a_passed_case():
    """Positive and negative control for the runner itself."""
    sheet = faults.build_fact_sheet(faults.load_snapshot(ROOT / "data" / "snapshots", "DGBF", "2026-08"))
    from app.llm import demo_draft
    base = demo_draft(set(sheet.facts))
    typed = {"id": "T-1", "family": "A1", "severity": "fatal",
             "edits": [{"at": "outlook.1", "en": "Returns were 5% last year.", "ja": "昨年の騰落率は5%でした。"}]}
    assert faults.run_case(typed, base, sheet)["status"] == "blocked"
    plain = {"id": "T-2", "family": "L", "severity": "accept", "edits": []}
    assert faults.run_case(plain, base, sheet)["status"] == "passed"


# ---- the script ----------------------------------------------------------------------------

def test_offline_script_writes_its_record(tmp_path, capsys):
    code = measure_offline.main(["--out", str(tmp_path)])
    assert code == 0
    record = json.loads((tmp_path / "offline_demo.json").read_text(encoding="utf-8"))
    assert {"m3", "m3_controls", "m4", "m5", "m15"} <= set(record)
    assert record["m3_controls"]["all_as_expected"]
    out = capsys.readouterr().out
    assert "M3 controls:" in out and "M5 legitimate:" in out
